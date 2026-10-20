#!/usr/bin/env python3
"""Genera la plantilla de casos dorados (.xlsx) para que quien sabe nómina capture
las cifras esperadas a mano. Las entradas salen de data/casos-dorados.mjs.

Uso:  node -e 'import("./data/casos-dorados.mjs").then(m=>console.log(JSON.stringify(m)))' > casos.json
      python3 scripts/plantilla-casos-dorados.py casos.json salida.xlsx
"""
import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

datos = json.load(open(sys.argv[1]))
casos, salidas = datos["CASOS_DORADOS"], datos["SALIDAS"]
salida_xlsx = sys.argv[2]

ARIAL = "Arial"
f_base = Font(name=ARIAL, size=10)
f_neg = Font(name=ARIAL, size=10, bold=True)
f_tit = Font(name=ARIAL, size=14, bold=True, color="0F3D3E")
f_cab = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
f_gris = Font(name=ARIAL, size=9, color="666666", italic=True)
cab = PatternFill("solid", fgColor="0F3D3E")
captura = PatternFill("solid", fgColor="FFF4C2")
grupo = PatternFill("solid", fgColor="E8F1F0")
borde = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
envolver = Alignment(wrap_text=True, vertical="top")

GRUPOS = {"ISR": "ISR y subsidio", "SUB": "ISR y subsidio", "EXE": "Exenciones (LISR art. 93)", "IMSS": "IMSS, RCV e INFONAVIT",
          "RCV": "IMSS, RCV e INFONAVIT", "INF": "IMSS, RCV e INFONAVIT", "ISN": "Impuesto sobre nóminas"}

wb = Workbook()

# ── Hoja 1: instrucciones ──
ins = wb.active
ins.title = "Instrucciones"
ins.column_dimensions["A"].width = 26
ins.column_dimensions["B"].width = 95
filas = [
    ("Casos dorados · kernel de nómina HRM", None),
    ("Para", "Rafael González (o quien calcule nómina a mano). Fecha de la plantilla: 19-oct-2026."),
    ("Qué es", "Cada renglón es un caso de nómina con datos completos. Calcule a mano (o en su hoja de siempre) la cifra que se pide y anótela. "
               "El sistema NO muestra su resultado aquí a propósito: su cifra es la referencia independiente contra la que se prueba el kernel."),
    ("Qué llenar", "Sólo las celdas amarillas de la hoja «Casos»: columna F (cifra esperada en pesos, con centavos), G (cómo la calculó o fundamento) "
                   "y H (observaciones). No cambie las demás columnas ni el orden de los renglones."),
    ("Formato de la cifra", "Número en pesos con dos decimales, sin signo de pesos ni separador de miles: 1222.87. Un resultado negativo va con signo menos: -150.25."),
    ("Si no aplica o no se puede", "Deje F vacía y explique en H (p. ej. «depende del criterio C1», «falta dato»). Un caso sin cifra no hace fallar la prueba."),
    ("Si un caso pide decidir", "Algunos renglones dicen ANOTE: además de la cifra, escriba en G qué método aplicó (por ejemplo, catorcenal diaria × 14 o proyección mensual). "
                                "Esa respuesta cierra un criterio de la bitácora (C1, C5, C13, C14)."),
    ("Supuestos comunes", "Ejercicio 2026. UMA $113.14 hasta el 31-ene y $117.31 desde el 1-feb. Salario mínimo general $315.04. Subsidio: 15.59% de la UMA mensual en enero "
                          "($536.21) y 15.02% desde febrero ($535.65); tope de ingresos $11,492.66 mensuales; divisor 30.4 (criterio C2). Tarifas del Anexo 8 RMF 2026. "
                          "Bases gravables ya netas de exenciones salvo en los casos de exenciones. Prima de riesgo indicada en cada caso. Trabajadores ficticios."),
    ("Ejemplo (ya resuelto)", "Caso de la revisión del maestro: semanal, base gravable $7,760.68, pago 6-ago-2026 → ISR a retener 1222.87 · "
                              "Cómo: renglón 6 semanal, 427.56 + (7,760.68 − 4,037.33) × 21.36% = 1,222.87; sin subsidio (base > tope semanal exacto $2,646.336)."),
    ("Al terminar", "Guarde el archivo con el mismo nombre y envíelo. El sistema lo importa (scripts/importar-casos-dorados.py) y compara cada cifra al centavo."),
]
for i, (k, v) in enumerate(filas, start=1):
    a = ins.cell(row=i, column=1, value=k)
    if v is None:
        a.font = f_tit
        ins.row_dimensions[i].height = 24
        continue
    a.font = f_neg
    a.alignment = envolver
    b = ins.cell(row=i, column=2, value=v)
    b.font = f_base
    b.alignment = envolver
    ins.row_dimensions[i].height = max(18, 15 * (len(v) // 95 + 1))
leyenda = len(filas) + 2
ins.cell(row=leyenda, column=1, value="Celda de captura").font = f_neg
c = ins.cell(row=leyenda, column=2, value="Amarillo = la llena usted")
c.fill = captura
c.font = f_base

# ── Hoja 2: casos ──
ws = wb.create_sheet("Casos")
cols = [("ID", 7), ("Regla", 17), ("Qué se prueba", 34), ("Datos del caso", 62), ("Importe que se pide", 30),
        ("Cifra esperada ($)", 16), ("Cómo lo calculó / fundamento", 44), ("Observaciones", 30)]
for j, (t, w) in enumerate(cols, start=1):
    h = ws.cell(row=1, column=j, value=t)
    h.font, h.fill, h.alignment, h.border = f_cab, cab, Alignment(wrap_text=True, vertical="center"), borde
    ws.column_dimensions[h.column_letter].width = w
ws.row_dimensions[1].height = 30
ws.freeze_panes = "C2"

dv = DataValidation(type="decimal", operator="between", formula1="-100000000", formula2="100000000", allow_blank=True,
                    error="Escriba la cifra en pesos con decimales, p. ej. 1222.87", errorTitle="Cifra no válida")
ws.add_data_validation(dv)

fila, ultimo_grupo = 2, None
for cs in casos:
    g = GRUPOS[cs["regla"].split("-")[1]]
    if g != ultimo_grupo:
        celda = ws.cell(row=fila, column=1, value=g)
        celda.font = f_neg
        for j in range(1, 9):
            ws.cell(row=fila, column=j).fill = grupo
        ultimo_grupo = g
        fila += 1
    valores = [cs["id"], cs["regla"], cs["titulo"], cs["legible"], salidas[cs["salida"]]["etiqueta"], None, None, None]
    for j, v in enumerate(valores, start=1):
        c = ws.cell(row=fila, column=j, value=v)
        c.font, c.alignment, c.border = f_base, envolver, borde
        if j >= 6:
            c.fill = captura
    ws.cell(row=fila, column=6).number_format = "#,##0.00;-#,##0.00"
    dv.add(ws.cell(row=fila, column=6))
    ws.row_dimensions[fila].height = max(28, 13 * (len(cs["legible"]) // 60 + 1))
    fila += 1

nota = ws.cell(row=fila + 1, column=1, value="Renglones: " + str(len(casos)) + " · casos: " + str(len({c['caso'] for c in casos})) +
               " · Generado de kernel/data/casos-dorados.mjs. Los encabezados de grupo (verde) no llevan cifra.")
nota.font = f_gris
ws.auto_filter.ref = "A1:H" + str(fila - 1)
for hoja in (ins, ws):
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = "1:1"
wb.save(salida_xlsx)
print("ok", salida_xlsx, len(casos), "renglones")
