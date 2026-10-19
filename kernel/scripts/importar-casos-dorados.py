#!/usr/bin/env python3
"""Importa la plantilla de casos dorados llenada por Rafael a data/casos-dorados-esperados.json.

Uso: python3 scripts/importar-casos-dorados.py casos-dorados.xlsx "Rafael González"
Lee la hoja «Casos»: ID (col A), cifra esperada en pesos (col F), cómo / fundamento (G), observaciones (H).
Convierte la cifra a centavos con redondeo medio hacia arriba y conserva lo ya importado de otros IDs.
"""
import json
import os
import sys
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook

archivo, quien = sys.argv[1], (sys.argv[2] if len(sys.argv) > 2 else "sin nombre")
destino = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "casos-dorados-esperados.json")
previo = json.load(open(destino)) if os.path.exists(destino) else {"casos": {}}

ws = load_workbook(archivo, data_only=True)["Casos"]
nuevos = 0
for fila in ws.iter_rows(min_row=2, values_only=True):
    ident, cifra, como, obs = fila[0], fila[5], fila[6], fila[7]
    if not ident or not str(ident).startswith("D"):
        continue
    if cifra is None or str(cifra).strip() == "":
        if obs or como:
            previo["casos"].setdefault(str(ident), {})
            previo["casos"][str(ident)].update({"esperado_c": None, "como": como, "observaciones": obs, "quien": quien, "fecha": date.today().isoformat()})
        continue
    texto = str(cifra).replace(",", "").replace("$", "").strip()
    c = int((Decimal(texto) * 100).quantize(Decimal(1), rounding=ROUND_HALF_UP))
    previo["casos"][str(ident)] = {"esperado_c": c, "como": como, "observaciones": obs, "quien": quien, "fecha": date.today().isoformat()}
    nuevos += 1

previo["descripcion"] = "Cifras esperadas de los casos dorados, capturadas a mano por quien sabe nómina (T1). No editar a mano: se importa de la plantilla."
json.dump(previo, open(destino, "w"), ensure_ascii=False, indent=1)
print("importados", nuevos, "→", destino)
